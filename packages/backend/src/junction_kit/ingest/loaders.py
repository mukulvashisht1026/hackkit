"""Load anything into text or a DataFrame: CSV/Excel, PDF, PPTX, web pages, plain text."""
from pathlib import Path

import duckdb
import httpx
import pandas as pd
from bs4 import BeautifulSoup


def load_table(path: str | Path) -> pd.DataFrame:
    p = Path(path)
    if p.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(p)
    if p.suffix.lower() == ".parquet":
        return pd.read_parquet(p)
    return pd.read_csv(p, encoding="utf-8", low_memory=False)


def sql(df_map: dict[str, pd.DataFrame], query: str) -> pd.DataFrame:
    """Run SQL over DataFrames with DuckDB, e.g. sql({'c': cultures}, 'select ward, count(*) from c group by 1')."""
    con = duckdb.connect()
    for name, df in df_map.items():
        con.register(name, df)
    return con.execute(query).df()


def pdf_to_pages(path: str | Path) -> list[str]:
    from pypdf import PdfReader

    return [(page.extract_text() or "").strip() for page in PdfReader(str(path)).pages]


def pptx_to_slides(path: str | Path) -> list[str]:
    from pptx import Presentation

    slides = []
    for slide in Presentation(str(path)).slides:
        parts = [s.text_frame.text for s in slide.shapes if getattr(s, "has_text_frame", False) and s.has_text_frame]
        slides.append("\n".join(p for p in parts if p.strip()))
    return slides


def html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "noscript", "svg"]):
        tag.decompose()
    text = soup.get_text("\n")
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def fetch_url(url: str, timeout: float = 20.0) -> str:
    """Plain fetch + clean text. For JS-heavy or bot-protected sites use partners.context_dev_scrape."""
    r = httpx.get(url, timeout=timeout, follow_redirects=True,
                  headers={"User-Agent": "Mozilla/5.0 (hackathon-bot)"})
    r.raise_for_status()
    return html_to_text(r.text)


def load_any_text(path: str | Path) -> list[str]:
    """Return a list of text chunks (pages/slides/the whole file) for any supported file."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        return pdf_to_pages(p)
    if ext == ".pptx":
        return pptx_to_slides(p)
    if ext in {".html", ".htm"}:
        return [html_to_text(p.read_text(errors="ignore"))]
    return [p.read_text(errors="ignore")]
